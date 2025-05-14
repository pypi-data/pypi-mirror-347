"""Main Module"""

import re
from pathlib import Path
from typing import Any, Dict, Optional, Union, Callable
from warnings import warn as warning
from functools import partial

from docx import Document
from docx.text.paragraph import Paragraph
from docx.text.hyperlink import Hyperlink
from docx.table import Table
from docx.oxml.ns import qn
from . import plugins
from .exceptions import RenderError
from .utils import fix_quotes, container_text_replace

PLUGINS = [plugins.image, plugins.table]

class DOCXRenderer:
    """DOCX Renderer class

    This class is used to render a DOCX template by replacing python statements
    with the result of evaluating the python statements.

    Attributes:
        template_path (str): Path to the DOCX template.
    """

    def __init__(self, template_path: Union[str, Path]):
        self.template_path = template_path
        self.document = Document(str(self.template_path))
        self.plugins = {}
        self.namespace = {}
        for plugin in PLUGINS:
            self.register_plugin(plugin.__name__, plugin)

    def register_plugin(self, name: str, func: Callable) -> None:
        """Register a plugin to be used in the template.

        Args:
            name (str): Name of the plugin.
            func (Callable): Function to be used as a plugin.

        Returns:
            None
        """
        self.plugins[name] = func
    
    def _substitute(self, item, skip_failed: bool = False) -> None:
        """Substitute the placeholders in the item with the result of evaluating
        the python statements.

        Args:
            item: The item to substitute. Can be a Paragraph or Table.
            skip_failed (bool, optional): Don't raise an error if some of the
                statements failed to render. Defaults to False.
        Returns:
            None
        """
        if isinstance(item, Paragraph) or isinstance(item, Hyperlink):
            # Hyperlinks are treated like paragraphs
            matches = re.finditer(r"{{{(.*?)}}}", item.text)
            for match in matches:
                parts = match.group(1).split(":", 1)
                try:
                    result = eval(fix_quotes(parts[0]), self.namespace)
                except Exception as ex:
                    if skip_failed:
                        continue
                    raise RenderError(f"Failed to evaluate '{parts[0]}'.") from ex
                if len(parts) > 1:
                    namespace = self.namespace.copy()
                    context = {
                        "result": result,
                        "paragraph": item,
                        "document": self.document,
                    }
                    for plugin_name, plugin in self.plugins.items():
                        func = partial(plugin, context)
                        namespace[plugin_name] = func 
                    try:
                        exec(fix_quotes(parts[1]), namespace)
                    except Exception as ex:
                        if skip_failed:
                            warning(
                                f"Failed to render {parts[0]}"
                            )
                            return
                        raise RenderError(
                            f"Failed to render {parts[0]}"
                        ) from ex
                else:
                    container_text_replace(item, match.group(0), str(result))
            if isinstance(item, Paragraph):
                for hyperlink in item.hyperlinks:
                    self._substitute(hyperlink, skip_failed)

        elif isinstance(item, Table):
            for row in item.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        self._substitute(para, skip_failed)
        else:
            return
        

    def render(
        self,
        output_path: Union[str, Path],
        methods_and_params: Optional[Dict[str, Any]] = None,
        skip_failed: bool = False,
    ) -> None:
        """Render DOCXRenderer template and save to output_path.

        Args:
            output_path (str): Path to the output DOCX file.
            methods_and_params (dict, optional): Dictionary of methods and parameters
                to be used in the template. Defaults to None.
            skip_failed (bool, optional): Don't raise an error if some of the
                statements failed to render. Defaults to False.

        Returns:
            None
        """
        if not Path(self.template_path).exists():
            raise FileNotFoundError(f"{self.template_path} not found")

        if not methods_and_params:
            methods_and_params = {}
        self.namespace.update(methods_and_params)

        for section in self.document.sections:
            for para in section.header.iter_inner_content():
                self._substitute(para, skip_failed)
            for para in section.footer.iter_inner_content():
                self._substitute(para, skip_failed)
            for item in section.iter_inner_content():
                self._substitute(item, skip_failed)
        
        for rel in self.document.part.rels.values():
            matches = re.finditer(r"%7b%7b%7b(.*?)%7d%7d%7d", rel.target_ref)
            for match in matches:
                matched_text = match.group(1)
                try:
                    result = eval(fix_quotes(matched_text), self.namespace)
                except Exception as ex:
                    if skip_failed:
                        continue
                    raise RenderError(f"Failed to evaluate '{rel.target_ref}'.") from ex
                if rel.is_external:
                    rel._target = rel.target_ref.replace(match.group(0), str(result))

        self.document.save(output_path)
